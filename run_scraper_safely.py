import os
import re
import json
import time
import random
import requests
from bs4 import BeautifulSoup
import openpyxl

EXCEL_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form.xlsx"
OUTPUT_EXCEL_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form_enriched.xlsx"
CACHE_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\scraped_cache.json"
JSON_OUTPUT_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\data.json"
LOG_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\scraper.log"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# Logger helper that writes to a UTF-8 file and flushes immediately
def log_message(msg):
    with open(LOG_PATH, "a", encoding="utf-8") as lf:
        lf.write(f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] {msg}\n")

# Clear log at start
with open(LOG_PATH, "w", encoding="utf-8") as lf:
    lf.write("--- Scraper Started ---\n")

# Load cache
if os.path.exists(CACHE_PATH):
    with open(CACHE_PATH, "r", encoding="utf-8") as f:
        cache = json.load(f)
    log_message(f"Loaded existing cache with {len(cache)} items.")
else:
    cache = {}
    log_message("No existing cache found. Initializing new cache.")

def save_cache():
    with open(CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

def extract_price_from_card(card):
    # 1. Offers JSON price
    offers_tag = card.find('offers')
    if offers_tag:
        offers_attr = offers_tag.get(':offers')
        if offers_attr:
            try:
                offers_json = json.loads(offers_attr)
                if offers_json:
                    price = offers_json[0].get('price')
                    if price:
                        return int(float(price))
            except Exception:
                pass
                
    # 2. Text containing "Розница X ₽"
    card_text = card.get_text(separator=' ')
    card_text = " ".join(card_text.split())
    match = re.search(r'Розница\s*([\d\s]+)\s*(?:₽|руб)', card_text)
    if match:
        price_str = match.group(1).replace(' ', '').replace('\xa0', '')
        try:
            return int(float(price_str))
        except ValueError:
            pass
            
    # 3. Text containing any price class with digits
    price_elem = card.find(class_=re.compile(r'price', re.I))
    if price_elem:
        price_text = " ".join(price_elem.text.split())
        match = re.search(r'([\d\s]+)\s*(?:₽|руб)', price_text)
        if match:
            price_str = match.group(1).replace(' ', '').replace('\xa0', '')
            try:
                return int(float(price_str))
            except ValueError:
                pass
                
    return None

def query_autoopt(search_term):
    url = f"https://www.autoopt.ru/search/index?search={search_term}"
    try:
        r = requests.get(url, headers=HEADERS, timeout=15)
        if r.status_code != 200:
            log_message(f"HTTP status {r.status_code} for search '{search_term}'")
            return None
        
        soup = BeautifulSoup(r.text, 'html.parser')
        cards = soup.find_all(lambda tag: tag.name == 'div' and tag.get('class') and 'n-catalog-item' in tag.get('class') and 'n-catalog-item__product' in tag.get('class'))
        
        results = []
        for card in cards:
            link_tag = card.find('a', href=re.compile(r'/catalog/\d+-.+'))
            if not link_tag:
                continue
            href = link_tag.get('href')
            title = link_tag.text.strip()
            
            price = extract_price_from_card(card)
            
            results.append({
                "title": title,
                "link": f"https://www.autoopt.ru{href}",
                "price": price
            })
        return results
    except Exception as e:
        log_message(f"Exception during request for '{search_term}': {e}")
        return None

def search_part(catalog_str):
    catalog_str = str(catalog_str).strip()
    if not catalog_str or catalog_str.lower() == 'nan':
        return None
        
    if catalog_str in cache:
        return cache[catalog_str]
        
    results = None
    
    # Strategy 1: Search exact string
    log_message(f"Searching exact: '{catalog_str}'")
    results = query_autoopt(catalog_str)
    time.sleep(random.uniform(0.3, 0.6))
    
    # Strategy 2: If slash exists, try splitting
    if (not results) and ('/' in catalog_str):
        parts = [p.strip() for p in catalog_str.split('/') if p.strip()]
        for part in reversed(parts):
            log_message(f"Searching split part: '{part}'")
            results = query_autoopt(part)
            time.sleep(random.uniform(0.3, 0.6))
            if results:
                break
                
    # Strategy 3: If spaces exist, try last word
    if (not results) and (' ' in catalog_str):
        words = [w.strip() for w in catalog_str.split() if w.strip()]
        last_word = words[-1]
        if any(c.isdigit() for c in last_word) and len(last_word) >= 4:
            log_message(f"Searching last word: '{last_word}'")
            results = query_autoopt(last_word)
            time.sleep(random.uniform(0.3, 0.6))
            
        # Strategy 4: Remove Cyrillic
        if not results:
            clean_words = [w for w in words if not re.search(r'[а-яА-Я]', w)]
            if clean_words:
                clean_str = " ".join(clean_words)
                log_message(f"Searching non-cyrillic: '{clean_str}'")
                results = query_autoopt(clean_str)
                time.sleep(random.uniform(0.3, 0.6))

    # Process and build data entry
    if results:
        best_match = None
        for res in results:
            if res['price'] is not None:
                best_match = res
                break
        if not best_match:
            best_match = results[0]
            
        data = {
            "title": best_match['title'],
            "link": best_match['link'],
            "price": best_match['price']
        }
        log_message(f"  Match Found: '{data['title']}' -> price: {data['price']} -> link: {data['link']}")
    else:
        data = {
            "title": "",
            "link": f"https://www.autoopt.ru/search/index?search={catalog_str}",
            "price": None
        }
        log_message("  No matches found.")
        
    cache[catalog_str] = data
    save_cache()
    return data

def main():
    log_message("Loading workbook...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    
    header_row = 11
    for r in range(1, 20):
        val = ws.cell(row=r, column=3).value
        if val and "Наименование" in str(val):
            header_row = r
            break
            
    log_message(f"Detected header row: {header_row}")
    ws.cell(row=header_row, column=7).value = "Ссылка на деталь"
    ws.cell(row=header_row+1, column=7).value = "6"
    
    total_processed = 0
    scraped_count = 0
    
    # Process Row 13 to ws.max_row
    for r in range(header_row + 2, ws.max_row + 1):
        num_val = ws.cell(row=r, column=2).value
        catalog_val = ws.cell(row=r, column=4).value
        
        if num_val is not None:
            try:
                # Validate numeric index row
                float(num_val)
                total_processed += 1
                
                if catalog_val:
                    # Retrieve details using safe search function that catches its own exceptions
                    try:
                        res = search_part(catalog_val)
                        if res:
                            if res['price'] is not None:
                                ws.cell(row=r, column=6).value = res['price']
                            ws.cell(row=r, column=7).value = res['link']
                            scraped_count += 1
                    except Exception as exc:
                        log_message(f"ERROR: Exception processing row {r} for catalog '{catalog_val}': {exc}")
                else:
                    ws.cell(row=r, column=7).value = ""
                    
            except ValueError:
                pass
                
        # Periodically save Excel
        if total_processed % 20 == 0:
            try:
                wb.save(OUTPUT_EXCEL_PATH)
                log_message(f"Saved progress to Excel. Processed {total_processed} rows...")
            except Exception as exc:
                log_message(f"ERROR: Failed to save Excel progress: {exc}")
                
    # Final Workbook save
    try:
        wb.save(OUTPUT_EXCEL_PATH)
        log_message(f"Final Excel saved successfully to {OUTPUT_EXCEL_PATH}")
    except Exception as exc:
        log_message(f"ERROR: Failed to save final Excel: {exc}")
        
    # Write JSON and JS data formats for the website
    log_message("Generating data files for website...")
    web_data = []
    
    # Reload final workbook to dump clean data
    wb_final = openpyxl.load_workbook(OUTPUT_EXCEL_PATH)
    ws_final = wb_final.active
    
    for r in range(header_row + 2, ws_final.max_row + 1):
        num_val = ws_final.cell(row=r, column=2).value
        if num_val is not None:
            try:
                float(num_val)
                name = ws_final.cell(row=r, column=3).value
                catalog = ws_final.cell(row=r, column=4).value
                unit = ws_final.cell(row=r, column=5).value
                price = ws_final.cell(row=r, column=6).value
                link = ws_final.cell(row=r, column=7).value
                
                web_data.append({
                    "id": int(float(num_val)),
                    "name": str(name).strip() if name else "",
                    "catalog": str(catalog).strip() if catalog else "",
                    "unit": str(unit).strip() if unit else "",
                    "price": int(price) if price is not None else None,
                    "link": str(link).strip() if link else ""
                })
            except ValueError:
                pass
                
    # Save to data.json
    try:
        with open(JSON_OUTPUT_PATH, "w", encoding="utf-8") as f:
            json.dump(web_data, f, ensure_ascii=False, indent=2)
        log_message(f"Saved data.json successfully to {JSON_OUTPUT_PATH}")
    except Exception as exc:
        log_message(f"ERROR: Failed to save data.json: {exc}")
        
    # Save to data.js for direct local CORS-free loading
    try:
        js_content = f"window.partsData = {json.dumps(web_data, ensure_ascii=False, indent=2)};"
        with open(r"c:\Users\User\Desktop\PYTHON TASKS\gleb\data.js", "w", encoding="utf-8") as f:
            f.write(js_content)
        log_message("Saved data.js successfully.")
    except Exception as exc:
        log_message(f"ERROR: Failed to save data.js: {exc}")
        
    log_message(f"Finished successfully. Processed {total_processed} rows, scraped {scraped_count} catalog items.")

if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        log_message(f"FATAL EXCEPTION in main execution: {e}")
