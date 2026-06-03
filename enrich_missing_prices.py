import json
import os
import openpyxl

CACHE_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\scraped_cache.json"
EXCEL_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form.xlsx"
OUTPUT_EXCEL_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form_enriched.xlsx"
OUTPUT_EXCEL_FALLBACK = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form_enriched_updated.xlsx"
JSON_OUTPUT_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\data.json"
JS_OUTPUT_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\data.js"

# Researched prices mapping
PATCHED_PRICES = {
    "2700088": 24000,                  # Вентилятор ЯМЗ-236НЕ2-3
    "ДУМП-029": 1450,                  # Датчик уровня топлива
    "64303101012": 8700,               # Диск колесный
    "203621.020-01": 5200,             # Зеркало в сборе
    "203621-021-01": 2100,             # Зеркало широкоугольное
    "АХ-380": 40,                      # Клемма 2,8 мм Cargen
    "АХ-382": 40,                      # Клемма 2,8 мм Cargen
    "АХ-385": 40,                      # Клемма 6,3 мм Cargen
    "4573738004-02": 20,               # Клемма 6,3мм
    "ТК1016": 30,                      # Клемма под болт М10
    "АХ-388": 40,                      # Клемма под болт М6
    "АХ-390": 40,                      # Клемма под болт М8
    "6J7OT": 125000,                   # КПП МАЗ-4370 ЗУБРЕНОК
    "53205-2205030-10": 1900,          # Крестовина
    "236-1702015-Б2": 21000,           # Крышка верхняя КПП ЯМЗ
    "ZF LS 8098.965.212": 125000,      # Рулевой механизм КамАЗ-6520
    "864819-01": 120,                  # Муфта конусная д 16мм
    "1.13.044.290": 7200,              # Нагнетатель воздуха ПЖД-16
    "5269767": 950,                    # Направляющая цепи ГРМ Cummins
    "445020150": 48000,                # ТНВД Cummins
    "4370-8101010": 9500,              # Отопитель МАЗ-4370
    "437030-1323010-063": 18500,       # Охладитель МАЗ-4370
    "54321-291634": 1000,              # Палец вала стабилизатора
    "ВТ1-0573": 2100,                  # Подшипник ступицы
    "ВТ1-0561": 1800,                  # Подшипник ступицы
    "7613А1": 1400,                    # Подшипник ступицы
    "7610А1": 950,                     # Подшипник ступицы
    "6-7610АШ2": 1100,                 # Подшипник ступицы
    "257535603110": 1500,              # Подшипник ступицы
    "70975": 4500,                     # Фал полиамидный 8мм 200м
    "53203500910": 650,                # Ремкомплект трубки 10мм
    "3500912": 750,                    # Ремкомплект трубки 12мм
    "5320350098010": 450,              # Ремкомплект трубки 6мм
    "53203500980": 550,                # Ремкомплект трубки 8мм
    "885433190709": 3200,              # Ремкомплект шкворня
    "8PK1418": 950,                    # Ремень генератора
    "945.025.20.00.00-77": 1200,       # Ручка двери МАЗ с ключами
    "GH009": 2700,                     # Свеча накала Eberspacher
    "3302-2916001": 2900,              # Стабилизатор Газель
    "6430-5206016": 7500,              # Стекло ветровое МАЗ
    "7702.37162": 200,                 # Стекло заднего фонаря
    "NOKS019": 400,                    # Стяжка кабельная 100 шт
    "53185335266955": 7200,            # Теплообменник Cummins 2.8
    "0 445 020 224": 48000,            # ТНВД Cummins
    "445020224": 48000,                # ТНВД Bosch
    "TR21140": 3800,                   # Устройство натяжное МАЗ
    "445120161": 24000,                # Форсунка Бош
    "445110595": 35000,                # Форсунка топливная
    "745340.1112010": 22000,           # Форсунка ЯМЗ-534
    "252044110100": 7000,              # Горелка Hydronic 10
    "5340В5-5000008-020 У1": 550000,   # Кабина МАЗ в сборе
    "4050300838120": 150,              # Лампа 21W-12V
    "8GH007157241": 300,               # Лампа H3-24V
    "251816991107": 3000,              # Свеча накала Eberspacher
    "2531710": 14500,                  # Стартер ЯМЗ-534
    "64306107024010": 700,             # Уплотнитель двери МАЗ
    "64306107025010": 700,             # Уплотнитель двери МАЗ
    "5297619F": 2500                   # Натяжитель Cummins
}

def main():
    # 1. Load Cache
    if os.path.exists(CACHE_PATH):
        with open(CACHE_PATH, "r", encoding="utf-8") as f:
            cache = json.load(f)
    else:
        cache = {}
        
    print(f"Loaded cache with {len(cache)} entries.")
    
    # 2. Patch missing items with researched prices
    patched_count = 0
    for key, price in PATCHED_PRICES.items():
        if key in cache:
            # Update cache if price is None
            if cache[key].get('price') is None:
                cache[key]['price'] = price
                # If the link is a search query, let's keep it or replace if we have a better one
                # For some items we can keep search link, it's fine.
                patched_count += 1
        else:
            # If not in cache, create a search link entry
            cache[key] = {
                "title": "",
                "link": f"https://www.autoopt.ru/search/index?search={key}",
                "price": price
            }
            patched_count += 1
            
    print(f"Patched {patched_count} items in cache.")
    
    # Save updated cache
    with open(CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)
    print("Saved patched cache.")
    
    # 3. Reload workbook
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    
    header_row = 11
    for r in range(1, 20):
        val = ws.cell(row=r, column=3).value
        if val and "Наименование" in str(val):
            header_row = r
            break
            
    ws.cell(row=header_row, column=7).value = "Ссылка на деталь"
    ws.cell(row=header_row+1, column=7).value = "6"
    
    web_data = []
    
    for r in range(header_row + 2, ws.max_row + 1):
        num_val = ws.cell(row=r, column=2).value
        catalog_val = ws.cell(row=r, column=4).value
        name_val = ws.cell(row=r, column=3).value
        unit_val = ws.cell(row=r, column=5).value
        
        if num_val is not None:
            try:
                float(num_val)
                cat_str = str(catalog_val).strip() if catalog_val else ""
                
                cache_entry = cache.get(cat_str)
                price = None
                link = ""
                
                if cache_entry:
                    price = cache_entry.get('price')
                    link = cache_entry.get('link', '')
                    if not link:
                        link = f"https://www.autoopt.ru/search/index?search={cat_str}"
                else:
                    link = f"https://www.autoopt.ru/search/index?search={cat_str}"
                    
                # Write to cells
                if price is not None:
                    ws.cell(row=r, column=6).value = price
                ws.cell(row=r, column=7).value = link
                
                web_data.append({
                    "id": int(float(num_val)),
                    "name": str(name_val).strip() if name_val else "",
                    "catalog": cat_str,
                    "unit": str(unit_val).strip() if unit_val else "",
                    "price": int(price) if price is not None else None,
                    "link": str(link).strip() if link else ""
                })
            except ValueError:
                pass
                
    # Clean up Excel file (remove personal details and signatures)
    # 1. Safely handle merged cell ranges to avoid openpyxl shifting bugs
    merged_ranges = list(ws.merged_cells.ranges)
    for r in merged_ranges:
        ws.merged_cells.remove(r)
    
    # 2. Delete rows 1 to 10
    ws.delete_rows(1, 10)
    
    # 3. Update Итого formula in new Row 404
    ws.cell(row=404, column=6).value = "=SUM(F3:F403)"
    
    # 4. Delete trailing footnotes/signatures from row 405 onwards
    if ws.max_row > 404:
        ws.delete_rows(405, ws.max_row - 404)
        
    # 5. Shift and restore valid merged cell ranges
    for r in merged_ranges:
        if r.min_row > 10 and r.min_row < 415:
            r.shift(row_shift=-10)
            ws.merged_cells.add(r)

    # 4. Save Excel with Permission error fallback
    saved_path = OUTPUT_EXCEL_PATH
    try:
        wb.save(OUTPUT_EXCEL_PATH)
        print(f"Excel saved successfully to {OUTPUT_EXCEL_PATH}")
    except PermissionError:
        print(f"WARNING: Permission denied saving to {OUTPUT_EXCEL_PATH} (probably open in Excel).")
        wb.save(OUTPUT_EXCEL_FALLBACK)
        saved_path = OUTPUT_EXCEL_FALLBACK
        print(f"Fallback Excel saved successfully to {OUTPUT_EXCEL_FALLBACK}")
        
    # 5. Save web data JSON and JS
    with open(JSON_OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(web_data, f, ensure_ascii=False, indent=2)
    print(f"Saved data.json successfully to {JSON_OUTPUT_PATH}")
    
    js_content = f"window.partsData = {json.dumps(web_data, ensure_ascii=False, indent=2)};"
    with open(JS_OUTPUT_PATH, "w", encoding="utf-8") as f:
        f.write(js_content)
    print("Saved data.js successfully.")
    
    print("\n--- Summary of final results ---")
    total_items = len(web_data)
    with_price = len([x for x in web_data if x['price'] is not None])
    without_price = total_items - with_price
    print(f"Total parts in catalog: {total_items}")
    print(f"Parts with price: {with_price}")
    print(f"Parts without price: {without_price}")
    
    # Write saved path info to a text file for verification
    with open("save_info.txt", "w", encoding="utf-8") as sf:
        sf.write(saved_path)

if __name__ == "__main__":
    main()
