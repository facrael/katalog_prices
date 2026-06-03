import os
import re
import json
import time
import random
import requests
from bs4 import BeautifulSoup
import openpyxl

# Configurations
EXCEL_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form.xlsx"
OUTPUT_EXCEL_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\form_enriched.xlsx"
CACHE_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\scraped_cache.json"
JSON_OUTPUT_PATH = r"c:\Users\User\Desktop\PYTHON TASKS\gleb\data.json"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# Load cache
if os.path.exists(CACHE_PATH):
    with open(CACHE_PATH, "r", encoding="utf-8") as f:
        cache = json.load(f)
else:
    cache = {}

def save_cache():
    with open(CACHE_PATH, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)

def extract_price_from_card(card):
    # Method 1: from offers tag
    offers_tag = card.find('offers')
    if offers_tag:
        offers_attr = offers_tag.get(':offers')
        if offers_attr:
            try:
                offers_json = json.loads(offers_attr)
                if offers_json:
                    price = offers_json[0].get('price')
                    if price:
                        return int(float(price))
            except Exception:
                pass
    
    # Method 2: search for Retail price in text
    card_text = card.get_text(separator=' ')
    card_text = " ".join(card_text.split())
    match = re.search(r'Розница\s*([\d\s]+)\s*(?:₽|руб)', card_text)
    if match:
        price_str = match.group(1).replace(' ', '').replace('\xa0', '')
        return int(float(price_str))
        
    # Method 3: look for any price class text
    price_elem = card.find(class_=re.compile(r'price', re.I))
    if price_elem:
        price_text = " ".join(price_elem.text.split())
        match = re.search(r'([\d\s]+)\s*(?:₽|руб)', price_text)
        if match:
            price_str = match.group(1).replace(' ', '').replace('\xa0', '')
            return int(float(price_str))
            
    return None

def query_autoopt(search_term):
    url = f"https://www.autoopt.ru/search/index?search={search_term}"
    try:
        r = requests.get(url, headers=HEADERS, timeout=10)
        if r.status_code != 200:
            return None
        
        soup = BeautifulSoup(r.text, 'html.parser')
        cards = soup.find_all(lambda tag: tag.name == 'div' and tag.get('class') and 'n-catalog-item' in tag.get('class') and 'n-catalog-item__product' in tag.get('class'))
        
        results = []
        for card in cards:
            link_tag = card.find('a', href=re.compile(r'/catalog/\d+-.+'))
            if not link_tag:
                continue
            href = link_tag.get('href')
            title = link_tag.text.strip()
            
            price = extract_price_from_card(card)
            
            results.append({
                "title": title,
                "link": f"https://www.autoopt.ru{href}",
                "price": price
            })
        return results
    except Exception as e:
        print(f"Network error searching for '{search_term}': {e}")
        return None

def search_part(catalog_str):
    catalog_str = str(catalog_str).strip()
    if not catalog_str or catalog_str.lower() == 'nan':
        return None
    
    # Check cache first
    if catalog_str in cache:
        return cache[catalog_str]
    
    # Search strategies
    # Strategy 1: Search exact string
    print(f"Searching exact: '{catalog_str}'...")
    results = query_autoopt(catalog_str)
    time.sleep(random.uniform(0.3, 0.7)) # Be polite to the server
    
    # Strategy 2: If slash exists, try splitting and searching the second part, then the first part
    if (not results) and ('/' in catalog_str):
        parts = [p.strip() for p in catalog_str.split('/') if p.strip()]
        for part in reversed(parts):
            print(f"Searching split part: '{part}'...")
            results = query_autoopt(part)
            time.sleep(random.uniform(0.3, 0.7))
            if results:
                break
                
    # Strategy 3: If spaces exist, try searching the last part if it looks like a code, or remove Cyrillic words
    if (not results) and (' ' in catalog_str):
        words = [w.strip() for w in catalog_str.split() if w.strip()]
        # Check if last word is a part number (contains digits)
        last_word = words[-1]
        if any(c.isdigit() for c in last_word) and len(last_word) >= 4:
            print(f"Searching last word: '{last_word}'...")
            results = query_autoopt(last_word)
            time.sleep(random.uniform(0.3, 0.7))
        
        # If still no results, try removing Cyrillic words and searching the rest
        if not results:
            clean_words = [w for w in words if not re.search(r'[а-яА-Я]', w)]
            if clean_words:
                clean_str = " ".join(clean_words)
                print(f"Searching non-cyrillic words: '{clean_str}'...")
                results = query_autoopt(clean_str)
                time.sleep(random.uniform(0.3, 0.7))

    # Process results
    if results:
        # We prefer the first result that has a price. If none have a price, take the first result.
        best_match = None
        for res in results:
            if res['price'] is not None:
                best_match = res
                break
        if not best_match:
            best_match = results[0]
            
        data = {
            "title": best_match['title'],
            "link": best_match['link'],
            "price": best_match['price']
        }
        print(f"  Found: {data['title']} -> {data['price']} ₽ ({data['link']})")
    else:
        # Fallback to search results page link, price is None
        data = {
            "title": "",
            "link": f"https://www.autoopt.ru/search/index?search={catalog_str}",
            "price": None
        }
        print("  Not found.")
        
    # Store in cache
    cache[catalog_str] = data
    save_cache()
    return data

def main():
    print("Loading workbook...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active
    
    # Set headers in Column 7 (G) for the links
    # Headers are at Row 10 (1-indexed) in our inspection
    # Let's verify which row has "Наименование запчасти или материала"
    header_row = 10 # Default from head JSON
    for r in range(1, 20):
        val = ws.cell(row=r, column=3).value
        if val and "Наименование" in str(val):
            header_row = r
            break
            
    print(f"Detected header row: {header_row}")
    ws.cell(row=header_row, column=7).value = "Ссылка на деталь"
    ws.cell(row=header_row+1, column=7).value = "6" # Number for the column header
    
    # Parse rows
    total_processed = 0
    scraped_count = 0
    
    # Loop from row header_row + 2 to the end
    # We find rows that have a numeric ID in column 2 (B)
    for r in range(header_row + 2, ws.max_row + 1):
        num_val = ws.cell(row=r, column=2).value
        catalog_val = ws.cell(row=r, column=4).value
        
        # Check if numeric
        if num_val is not None:
            try:
                # Try converting to float to check if it's a part row
                float(num_val)
                total_processed += 1
                
                # Retrieve price and link
                if catalog_val:
                    res = search_part(catalog_val)
                    if res:
                        # Write price in column 6 (F)
                        if res['price'] is not None:
                            ws.cell(row=r, column=6).value = res['price']
                        
                        # Write link in column 7 (G)
                        ws.cell(row=r, column=7).value = res['link']
                        scraped_count += 1
                else:
                    # No catalog number
                    ws.cell(row=r, column=7).value = ""
                    
            except ValueError:
                # Not a numeric data row
                pass
                
        # Periodically save Excel file and cache
        if total_processed % 20 == 0:
            wb.save(OUTPUT_EXCEL_PATH)
            print(f"Saved progress to Excel. Processed {total_processed} rows...")
            
    # Final saves
    wb.save(OUTPUT_EXCEL_PATH)
    print("Excel saved to", OUTPUT_EXCEL_PATH)
    
    # Save cache
    save_cache()
    
    # Generate data.json for the website
    print("Generating website data.json...")
    web_data = []
    # Read the final sheet to extract clean data
    for r in range(header_row + 2, ws.max_row + 1):
        num_val = ws.cell(row=r, column=2).value
        if num_val is not None:
            try:
                float(num_val)
                name = ws.cell(row=r, column=3).value
                catalog = ws.cell(row=r, column=4).value
                unit = ws.cell(row=r, column=5).value
                price = ws.cell(row=r, column=6).value
                link = ws.cell(row=r, column=7).value
                
                web_data.append({
                    "id": int(float(num_val)),
                    "name": str(name).strip() if name else "",
                    "catalog": str(catalog).strip() if catalog else "",
                    "unit": str(unit).strip() if unit else "",
                    "price": int(price) if price is not None else None,
                    "link": str(link).strip() if link else ""
                })
            except ValueError:
                pass
                
    with open(JSON_OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(web_data, f, ensure_ascii=False, indent=2)
    print("JSON saved to", JSON_OUTPUT_PATH)
    print(f"Done! Processed {total_processed} rows, scraped {scraped_count} catalog items.")

if __name__ == "__main__":
    main()
